# RLR · La Vela — reconstruye el modelo de negocio v2 a partir del v1 — Ricardo López Reyero
# Conserva las 5,525 fórmulas del v1 y encima:
#   1) lo viste con la guía de estilos (blanco y negro, docs/14),
#   2) agrega el cartucho retornable (supuestos, hoja «Cartucho» y selector en el costeo),
#   3) agrega la hoja «Distribuidor» (cuánto gana un distribuidor según su tamaño),
#   4) rehace la gráfica del Resumen en grises.
# Uso: python3 herramientas/construir_modelo.py   (luego recalcular con LibreOffice)
import copy
import sys
from pathlib import Path

import openpyxl
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName

_RLR = "Ricardo López Reyero"; _k = "EYE"; _rev = 181218  # RLR

RAIZ = Path(__file__).resolve().parent.parent
ORIGEN = RAIZ / "modelo" / "Modelo_Negocio_Velas_Rinde.xlsx"
DESTINO = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / "modelo" / "La_Vela_Modelo_de_Negocio_v2.xlsx"

# ── Guía de estilos (docs/14): negro, blanco y grises ──
NEGRO, TINTA, G700, G500, G300, G100, BLANCO = "000000", "111111", "444444", "6B6B6B", "C8C8C8", "F3F3F3", "FFFFFF"
relleno = lambda c: PatternFill("solid", fgColor=c)
DELGADO = Side(style="thin", color=G500)
GRUESO = Side(style="medium", color=NEGRO)
MARCO = Border(left=DELGADO, right=DELGADO, top=DELGADO, bottom=DELGADO)        # dato que se puede cambiar
MARCO_CLAVE = Border(left=GRUESO, right=GRUESO, top=GRUESO, bottom=GRUESO)      # selector o supuesto clave
F_PESOS = '\\$#,##0.00;"($"#,##0.00\\);\\-'
F_ENTERO = '#,##0;(#,##0);\\-'
F_PESOS0 = '\\$#,##0;"($"#,##0\\);\\-'
F_PCT = '0.0%;\\(0.0%\\);\\-'


def fuente(c, **k):
    base = dict(name=c.font.name or "Arial", sz=c.font.sz or 10, b=c.font.b, i=c.font.i, color=NEGRO)
    base.update(k)
    return Font(**base)


# RLR · traduce el formato de color del v1 a blanco y negro, celda por celda
def vestir(c):
    f = c.font
    col = f.color.rgb if f.color is not None and f.color.type == "rgb" else None
    fondo = c.fill.fgColor.rgb if c.fill.fill_type else None
    azul, verde = col == "FF0000FF", col == "FF008000"
    if col == "FF1F3A5F":  # títulos
        c.font = Font(name="Georgia", sz=18, color=NEGRO) if (f.sz or 0) >= 14 else Font(name="Arial", sz=11, b=True, color=NEGRO)
        return
    if fondo == "FF1F3A5F":  # encabezado de tabla
        c.font = fuente(c, b=True, color=BLANCO); c.fill = relleno(NEGRO); return
    if azul and fondo == "FFFFFF00":  # supuesto clave
        c.font = fuente(c, b=True); c.fill = PatternFill(fill_type=None); c.border = MARCO_CLAVE; return
    if azul and fondo == "FFDCE6F1":  # ejes de las tablas de sensibilidad
        c.font = fuente(c, b=True); c.fill = relleno(G300); c.border = MARCO; return
    if azul:  # dato editable
        c.font = fuente(c); c.border = MARCO; return
    if fondo in ("FFE2EFDA", "FFFFFF00"):  # resultado clave
        c.font = fuente(c, b=True, color=BLANCO); c.fill = relleno(TINTA); return
    if fondo == "FFDCE6F1":
        c.font = fuente(c); c.fill = relleno(G300); return
    if fondo == "FFF2F2F2":
        c.font = fuente(c); c.fill = relleno(G100); return
    if verde:
        c.font = fuente(c); return
    if col == "FF595959":
        c.font = fuente(c, color=G500); return
    c.font = fuente(c)


wb = openpyxl.load_workbook(ORIGEN)
for ws in wb:
    for fila in ws.iter_rows():
        for c in fila:
            if c.value is not None or c.fill.fill_type:
                vestir(c)
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)

# ── Estilos para lo nuevo ──
def titulo(ws, texto, sub=None):
    ws["A1"] = texto; ws["A1"].font = Font(name="Georgia", sz=18, color=NEGRO); ws.row_dimensions[1].height = 26
    if sub:
        ws["A2"] = sub; ws["A2"].font = Font(name="Arial", sz=9, color=G500)

def seccion(ws, fila, texto, hasta):
    for col in range(1, hasta + 1):
        c = ws.cell(fila, col); c.fill = relleno(G300); c.font = Font(name="Arial", sz=10, b=True, color=NEGRO)
    ws.cell(fila, 1).value = texto

def encabezado(ws, fila, textos):
    for i, t in enumerate(textos, 1):
        c = ws.cell(fila, i, t); c.fill = relleno(NEGRO); c.font = Font(name="Arial", sz=10, b=True, color=BLANCO)
        c.alignment = Alignment(horizontal="left" if i == 1 else "right", vertical="center", wrap_text=True)

def celda(ws, ref, valor, fmt=None, negrita=False, entrada=False, clave=False, nota=False, resultado=False):
    c = ws[ref]; c.value = valor
    c.font = Font(name="Arial", sz=9 if nota else 10, b=negrita or clave or resultado, color=BLANCO if resultado else (G500 if nota else NEGRO))
    if fmt: c.number_format = fmt
    if entrada: c.border = MARCO
    if clave: c.border = MARCO_CLAVE
    if resultado: c.fill = relleno(TINTA)
    if nota: c.alignment = Alignment(wrap_text=True, vertical="top")
    return c

def nombre(clave, hoja, ref):
    wb.defined_names[clave] = DefinedName(clave, attr_text=f"{hoja}!{ref}")

# ── 1. Nombre del producto y leyendas (ya no hay colores) ──
for hoja, ref in (("Supuestos", "B19"), ("Supuestos", "B20"), ("Costeo", "C4"), ("Costeo", "D4")):
    wb[hoja][ref].value = wb[hoja][ref].value.replace(" Rinde+", "")
ini = wb["Inicio"]
ini["A1"] = "Modelo de negocio · La Vela"
ini["A2"] = ("Veladora diseñada para durar, con cartucho retornable, distribuida en tiendas y misceláneas a través de "
             "distribuidores con rutas, más canales directos y línea premium. Modelo a 36 meses en pesos mexicanos. Versión 2 · octubre de 2026.")
ini["A6"] = "2. Cambia cualquier celda con marco. Todo se recalcula solo. Las de marco grueso son las que más mueven el resultado."
ini["A8"] = ini["A8"].value  # sin cambio
for f in range(9, 40):
    for col in (1, 2):
        c = ini.cell(f, col); c.value = None; c.border = Border(); c.fill = PatternFill(fill_type=None)
celda(ini, "A9", "5. En 'Supuestos', celda D184: elige 1 = repuesto con funda o 2 = cartucho retornable. La hoja 'Cartucho' explica la diferencia.")
ini["A11"] = "Cómo se lee"; ini["A11"].font = Font(name="Arial", sz=11, b=True)
for f, (a, b) in enumerate([("Marco delgado", "Dato que puedes cambiar."), ("Marco grueso", "Selector o supuesto clave: lo que más mueve el resultado."),
                            ("Fondo negro", "Resultado clave, o encabezado de tabla."), ("Sin marco", "Fórmula: no tocar.")], 12):
    celda(ini, f"A{f}", a); celda(ini, f"B{f}", b)
ini["A12"].border = MARCO; ini["A13"].border = MARCO_CLAVE; ini["A13"].font = Font(name="Arial", sz=10, b=True)
ini["A14"].fill = relleno(TINTA); ini["A14"].font = Font(name="Arial", sz=10, b=True, color=BLANCO)
ini["A17"] = "Hojas"; ini["A17"].font = Font(name="Arial", sz=11, b=True)
HOJAS = [
    ("Resumen", "Resultados por año, capital, punto de equilibrio y comparación Base vs Optimizado."),
    ("Supuestos", "Todos los datos editables con su fuente. Selectores de escenario."),
    ("Costeo", "Costo por pieza de cada producto, de la cera a la caja."),
    ("Cartucho", "NUEVA. El cartucho retornable: cuánto cuesta cada llenado, cuántos hay que comprar y qué cambia contra la funda."),
    ("Cascada", "Quién gana qué en cada pieza, a dónde se va cada peso y economía de canales directos."),
    ("Todos_ganan", "Lo que gana el cliente, la tienda, el distribuidor y nosotros."),
    ("Distribuidor", "NUEVA. Cuánto gana un distribuidor según cuántas tiendas surte, y cómo se califica una solicitud."),
    ("Volumen", "Tiendas, piezas y canales mes a mes."),
    ("Resultados", "Estado de resultados mensual y anual."),
    ("Flujo", "Flujo de efectivo, capital de trabajo, necesidad de capital y recuperación."),
    ("Capex_Capacidad", "Equipo por nivel, costo de transformación propia y capacidad de planta."),
    ("Palancas", "Cuánto dinero deja cada mejora de eficiencia, e ideas extra."),
    ("Sensibilidad", "Tablas de qué pasa si cambian parafina, precio, márgenes, volumen y vaso."),
    ("Competencia", "Precios de competencia y captura de pruebas de horas reales."),
    ("Prueba_encendido", "Plantilla para registrar las pruebas de encendido."),
    ("Riesgos", "Supuestos por validar con responsable y fecha."),
    ("Fuentes", "De dónde salió cada dato."),
]
ini._hyperlinks = []  # las ligas del v1 apuntaban a las filas viejas
for f in range(1, 40):
    ini.cell(f, 1).hyperlink = None
for f, (a, b) in enumerate(HOJAS, 18):
    celda(ini, f"A{f}", a, negrita=True); celda(ini, f"B{f}", b)
    ini[f"A{f}"].hyperlink = f"#'{a}'!A1"
    ini[f"A{f}"].font = Font(name="Arial", sz=10, b=True, color=NEGRO, u="single")
for f, t in ((36, "Página del modelo: https://vela.capitaltorreon.com/modelo"), (37, "Ing. Ricardo López Reyero · Torreón, Coahuila")):
    celda(ini, f"A{f}", t, nota=True); ini[f"A{f}"].alignment = Alignment(wrap_text=False)

sup = wb["Supuestos"]
sup["A2"] = ("Marco delgado = dato que puedes cambiar · Marco grueso = selector o supuesto clave · La columna F (Activo) es la que usa "
             "todo el modelo. Todos los precios al público incluyen IVA; los costos de insumos son sin IVA.")
sen = wb["Sensibilidad"]
sen["A2"] = "Fondo negro = perdemos. Todo lo demás se queda como está en el escenario activo."
wb["Prueba_encendido"]["A2"] = wb["Prueba_encendido"]["A2"].value.replace("Celdas azules = captura", "Celdas con marco = captura")
# Sensibilidad: una sola regla, en blanco y negro
rangos = [str(cf.sqref) for cf in sen.conditional_formatting]
sen.conditional_formatting = openpyxl.formatting.formatting.ConditionalFormattingList()
for r in dict.fromkeys(rangos):
    sen.conditional_formatting.add(r, CellIsRule(operator="lessThan", formula=["0"], fill=relleno(TINTA), font=Font(color=BLANCO, b=True)))

# ── 2. Supuestos del cartucho retornable (docs/08) ──
F = 183
seccion(sup, F, "17. CARTUCHO RETORNABLE (el repuesto que regresa a la tienda)", 7)
filas = [
    # clave, concepto, unidad, base, optimizado, formato, nota
    ("cart_costo", "Costo de un cartucho de aluminio nuevo", "$ / pza", 3.2, 2.6, F_PESOS, "ESTIMADO: cotizar por millar con un proveedor de estampado o troquel."),
    ("cart_vueltas", "Vueltas que aguanta un cartucho antes de reciclarse", "vueltas", 8, 12, F_ENTERO, "ESTIMADO: medir en prototipo cuántas veces se rellena sin deformarse."),
    ("cart_retorno", "Cartuchos que el cliente regresa a la tienda", "%", 0.7, 0.8, F_PCT, "ESTIMADO: medir en el piloto de 50 tiendas. Es el supuesto que más pesa."),
    ("cart_deposito", "Depósito que paga el cliente por cartucho (con IVA)", "$ / pza", 5, 5, F_PESOS, "Como el envase de refresco retornable. Se descuenta al regresar el vacío."),
    ("cart_reproceso", "Reproceso por vuelta (calentar, sacar el cabo, revisar)", "$ / pza", 0.5, 0.35, F_PESOS, "ESTIMADO. No se lava: el residuo se integra a la cera nueva."),
    ("cart_log_inv", "Logística inversa por cartucho (caja de 24 en la ruta)", "$ / pza", 0.1, 0.05, F_PESOS, "ESTIMADO. Regresa en la misma visita semanal del distribuidor."),
    ("cart_ciclo", "Semanas que tarda un cartucho en dar la vuelta completa", "semanas", 5, 4, F_ENTERO, "Casa (1) + tienda + ruta + planta. Define cuántos cartuchos hay en circulación."),
]
celda(sup, f"A{F+1}", "sel_rep"); celda(sup, f"B{F+1}", "FORMATO DEL REPUESTO (1 = con funda, como en la versión 1 · 2 = cartucho retornable)", negrita=True)
celda(sup, f"C{F+1}", "1 / 2"); celda(sup, f"D{F+1}", 1, clave=True)
celda(sup, f"F{F+1}", f'=IF(D{F+1}=1,"Repuesto con funda","Cartucho retornable")', negrita=True)
celda(sup, f"G{F+1}", "Con 2, el costeo cambia la funda por el costo neto del cartucho (hoja Cartucho) y todo el modelo se recalcula.", nota=True)
nombre("sel_rep", "Supuestos", f"$D${F+1}")
for i, (clave, concepto, unidad, base, opt, fmt, nota) in enumerate(filas, F + 2):
    celda(sup, f"A{i}", clave); celda(sup, f"B{i}", concepto); celda(sup, f"C{i}", unidad)
    celda(sup, f"D{i}", base, fmt, entrada=True); celda(sup, f"E{i}", opt, fmt, entrada=True)
    celda(sup, f"F{i}", f"=IF(sel_costos=1,D{i},E{i})", fmt)
    celda(sup, f"G{i}", nota, nota=True); sup[f"G{i}"].alignment = Alignment(wrap_text=False)
    sup[f"A{i}"].font = copy.copy(sup["A19"].font); sup[f"C{i}"].font = copy.copy(sup["C19"].font)
    nombre(clave, "Supuestos", f"$F${i}")
sup[f"A{F+1}"].font = copy.copy(sup["A19"].font)
sup[f"G{F+1}"].alignment = Alignment(wrap_text=False)

# ── 3. Hoja Cartucho ──
ca = wb.create_sheet("Cartucho", index=wb.sheetnames.index("Costeo") + 1)
ca.sheet_view.showGridLines = False
titulo(ca, "El cartucho retornable", "El repuesto es una copa de aluminio con la cera y la mecha ya puestas. Regresa vacío a la tienda, se rellena y vuelve a salir. Los datos vienen de Supuestos, sección 17.")
for col, ancho in zip("ABCD", (62, 16, 16, 78)):
    ca.column_dimensions[col].width = ancho

seccion(ca, 4, "EL CICLO", 4)
for i, t in enumerate([
    "1. El cliente mete el cartucho completo al vaso y lo prende. La cera arde dentro del aluminio; el vaso no se ensucia.",
    "2. Cuando se acaba, regresa el cartucho vacío a la misma tienda y le descuentan el depósito en el siguiente.",
    "3. La tienda junta los vacíos en una caja de 24. La ruta del distribuidor se los lleva en su visita de siempre.",
    "4. En planta no se lava: se calienta, se saca el cabo, se pone mecha nueva y se rellena. El residuo se integra a la cera nueva.",
    "5. Cuando un cartucho ya no sirve, el aluminio se recicla.",
], 5):
    celda(ca, f"A{i}", t)

seccion(ca, 11, "CUÁNTO CUESTA CADA LLENADO", 4)
encabezado(ca, 12, ["Concepto", "Valor", "Unidad", "Cómo se calcula"])
datos = [
    (13, "Costo de un cartucho nuevo", "=cart_costo", F_PESOS, "$ / pza", "Supuestos."),
    (14, "Cartuchos que regresan", "=cart_retorno", F_PCT, "%", "Supuestos. De cada 100 cartuchos vendidos, cuántos vuelven a la tienda."),
    (15, "Vueltas máximas por cartucho", "=cart_vueltas", F_ENTERO, "vueltas", "Supuestos."),
    (16, "Llenados que da un cartucho en su vida", "=IF(cart_retorno>=1,cart_vueltas,(1-cart_retorno^cart_vueltas)/(1-cart_retorno))", "0.00", "llenados",
     "En cada vuelta se pierde el que no regresa. Con 80% de retorno y 12 vueltas, un cartucho se llena 4.7 veces en promedio."),
    (17, "Aluminio por llenado", "=B13/B16", F_PESOS, "$ / llenado", "Costo del cartucho repartido entre sus llenados."),
    (18, "Reproceso y logística inversa por llenado", "=cart_retorno*(cart_reproceso+cart_log_inv)", F_PESOS, "$ / llenado", "Solo aplica a los que regresan."),
    (19, "Depósito que nadie reclama, sin IVA", "=(1-cart_retorno)*cart_deposito/(1+iva)", F_PESOS, "$ / llenado",
     "El cliente que no regresa el cartucho deja su depósito. Supone que ese depósito llega hasta nosotros: VALIDAR con el distribuidor."),
]
for f, concepto, formula, fmt, unidad, nota in datos:
    celda(ca, f"A{f}", concepto); celda(ca, f"B{f}", formula, fmt); celda(ca, f"C{f}", unidad); celda(ca, f"D{f}", nota, nota=True)
celda(ca, "A20", "COSTO NETO DEL CARTUCHO POR LLENADO", resultado=True); celda(ca, "B20", "=B17+B18-B19", F_PESOS, resultado=True)
celda(ca, "C20", "$ / llenado", resultado=True); celda(ca, "D20", "Aluminio + reproceso − depósitos no reclamados. Es lo que entra al costeo cuando el selector está en 2.", nota=True)
nombre("cart_neto", "Cartucho", "$B$20")
celda(ca, "A21", "Costo sin contar los depósitos (lectura conservadora)"); celda(ca, "B21", "=B17+B18", F_PESOS); celda(ca, "C21", "$ / llenado")
celda(ca, "D21", "Si el depósito no reclamado se queda en la tienda o en el distribuidor.", nota=True)
celda(ca, "A22", "Funda del repuesto (versión 1)"); celda(ca, "B22", "=funda_rep", F_PESOS); celda(ca, "C22", "$ / pza"); celda(ca, "D22", "Lo que hoy usa el modelo con el selector en 1.", nota=True)
celda(ca, "A23", "Ahorro por repuesto contra la funda", negrita=True); celda(ca, "B23", "=B22-B20", F_PESOS, negrita=True); celda(ca, "C23", "$ / pza")
celda(ca, "D23", "Positivo = el cartucho retornable sale más barato que la funda.", nota=True)

seccion(ca, 25, "QUÉ CAMBIA EN UN AÑO (año 2 del escenario activo)", 4)
encabezado(ca, 26, ["Concepto", "Valor", "Unidad", "Cómo se calcula"])
datos = [
    (27, "Velas con vaso al año", "=Palancas!B6", F_ENTERO, "pzas", "Cada una sale con su primer cartucho."),
    (28, "Repuestos al año", "=Palancas!B7", F_ENTERO, "pzas", "Cada uno es un llenado."),
    (29, "Llenados al año", "=B27+B28", F_ENTERO, "llenados", ""),
    (30, "Cartuchos nuevos que hay que comprar al año", "=B29/B16", F_ENTERO, "pzas", "Llenados ÷ llenados por cartucho."),
    (31, "Compra anual de cartuchos", "=B30*B13", F_PESOS0, "$", ""),
    (32, "Reproceso y logística inversa al año", "=B29*B18", F_PESOS0, "$", ""),
    (33, "Depósitos no reclamados al año, sin IVA", "=B29*B19", F_PESOS0, "$", "Ingreso que compensa los cartuchos perdidos."),
    (34, "Costo neto anual del sistema de cartucho", "=B31+B32-B33", F_PESOS0, "$", ""),
    (35, "Costo anual de las fundas (versión 1)", "=B28*B22", F_PESOS0, "$", "Solo los repuestos llevan funda."),
]
for f, concepto, formula, fmt, unidad, nota in datos:
    celda(ca, f"A{f}", concepto); celda(ca, f"B{f}", formula, fmt); celda(ca, f"C{f}", unidad); celda(ca, f"D{f}", nota, nota=True)
celda(ca, "A36", "DIFERENCIA AL AÑO: CARTUCHO RETORNABLE CONTRA FUNDA", resultado=True); celda(ca, "B36", "=B35-B34", F_PESOS0, resultado=True)
celda(ca, "C36", "$", resultado=True); celda(ca, "D36", "Positivo = el cartucho deja más dinero. No incluye la merma (1.5%) ni el vidrio que se ahorra en ruta.", nota=True)

seccion(ca, 38, "CARTUCHOS EN CIRCULACIÓN (año 2)", 4)
encabezado(ca, 39, ["Concepto", "Valor", "Unidad", "Cómo se calcula"])
datos = [
    (40, "Llenados al mes", "=B29/12", F_ENTERO, "llenados", ""),
    (41, "Semanas que tarda la vuelta completa", "=cart_ciclo", F_ENTERO, "semanas", "Supuestos."),
    (42, "Cartuchos en circulación", "=B40*B41/semanas_mes", F_ENTERO, "pzas", "Los que hay en casas, tiendas, ruta y planta al mismo tiempo."),
    (43, "Inversión en la flotilla de cartuchos", "=B42*B13", F_PESOS0, "$", "Se compra una vez, al arrancar. No está en la hoja Flujo."),
    (44, "Depósitos cobrados por los cartuchos en circulación", "=B42*cart_deposito", F_PESOS0, "$", "Es dinero del cliente: se le debe hasta que regrese el cartucho."),
    (45, "Cartuchos que se dan de baja al año", "=B30", F_ENTERO, "pzas", "Perdidos o reciclados; son los que se reponen."),
]
for f, concepto, formula, fmt, unidad, nota in datos:
    celda(ca, f"A{f}", concepto); celda(ca, f"B{f}", formula, fmt); celda(ca, f"C{f}", unidad); celda(ca, f"D{f}", nota, nota=True)

seccion(ca, 47, "LO QUE HAY QUE VALIDAR", 4)
for i, t in enumerate([
    "Cuántas vueltas aguanta un cartucho antes de deformarse (prototipo).",
    "Qué porcentaje de los clientes regresa el cartucho, por tienda (piloto de 50 tiendas, 4 semanas).",
    "El costo del cartucho por millar (troquel) y el costo de reproceso por vuelta.",
    "La mecha calibrada para el diámetro del cartucho (~5.8 cm).",
    "Quién se queda con el depósito no reclamado: la tienda, el distribuidor o nosotros.",
    "Que el borde de aluminio se vea bien dentro del vaso.",
], 48):
    celda(ca, f"A{i}", "· " + t)
ca.freeze_panes = "A4"

# El costeo usa el cartucho cuando el selector está en 2
co = wb["Costeo"]
co["A17"] = "Funda del repuesto o cartucho retornable"
for col in "CEFG":
    co[f"{col}17"] = "=IF(sel_rep=1,0,cart_neto)"
co["D17"] = "=IF(sel_rep=1,funda_rep,cart_neto)"
co["H17"] = "Supuestos D184: 1 = funda (versión 1) · 2 = cartucho retornable, en todas las piezas (ver hoja Cartucho)."
co["H17"].font = Font(name="Arial", sz=9, color=G500)

# ── 4. Hoja Distribuidor ──
di = wb.create_sheet("Distribuidor", index=wb.sheetnames.index("Todos_ganan") + 1)
di.sheet_view.showGridLines = False
titulo(di, "Lo que gana un distribuidor", "Para la llamada con cada distribuidor. Cambia las celdas con marco: tiendas que surte y piezas por tienda. Lo demás sale del escenario activo.")
di.column_dimensions["A"].width = 56
for col in "BCDE":
    di.column_dimensions[col].width = 18
di.column_dimensions["F"].width = 60
seccion(di, 4, "POR TAMAÑO DE DISTRIBUIDOR", 6)
encabezado(di, 5, ["Concepto", "C · Chico", "B · Mediano", "A · Grande", "Tu caso", "Cómo se calcula"])
entradas = {"B": (100, 5), "C": (300, 8), "D": (1000, 8), "E": (500, 8)}
celda(di, "A6", "Tiendas que surte"); celda(di, "A7", "Piezas que vende cada tienda por semana")
celda(di, "F6", "Dato del distribuidor.", nota=True); celda(di, "F7", "ESTIMADO: se mide en el piloto. El escenario Medio usa 8.", nota=True)
renglones = [
    (8, "Piezas al mes", "={c}6*{c}7*semanas_mes", F_ENTERO, "Tiendas × piezas × 4.33 semanas.", False),
    (9, "El distribuidor gana por pieza", "=Cascada!$B$32", F_PESOS, "Promedio ponderado por la mezcla de productos (hoja Cascada).", False),
    (10, "EL DISTRIBUIDOR GANA AL MES", "={c}8*{c}9", F_PESOS0, "", True),
    (11, "El distribuidor gana al año", "={c}10*12", F_PESOS0, "Sin contar la estacionalidad.", False),
    (12, "Cada tienda gana al mes", "={c}7*semanas_mes*Cascada!$B$31", F_PESOS0, "Piezas de la tienda × lo que gana por pieza.", False),
    (13, "Su compra mensual a nosotros, con IVA", "={c}8*Cascada!$B$30*(1-m_tienda)*(1-m_dist)", F_PESOS0, "Lo que nos paga el distribuidor cada mes.", False),
    (14, "Repuestos al mes (cartuchos que van y vienen)", "={c}8*mix_rep", F_ENTERO, "Piezas × parte de la mezcla que es repuesto.", False),
    (15, "Nosotros: utilidad bruta al mes con este distribuidor", "={c}8*Cascada!$B$28", F_PESOS0, "Antes de gastos fijos.", False),
]
for col, (tiendas, piezas) in entradas.items():
    celda(di, f"{col}6", tiendas, F_ENTERO, entrada=True); celda(di, f"{col}7", piezas, F_ENTERO, entrada=True)
for f, concepto, formula, fmt, nota, clave in renglones:
    celda(di, f"A{f}", concepto, resultado=clave)
    for col in entradas:
        celda(di, f"{col}{f}", formula.format(c=col), fmt, resultado=clave)
    celda(di, f"F{f}", nota, nota=True)

seccion(di, 17, "CÓMO SE CALIFICA UNA SOLICITUD (interno, no se publica)", 6)
encabezado(di, 18, ["Pregunta del formulario", "1 punto", "2 puntos", "3 puntos", "", ""])
for f, fila in enumerate([
    ("Puntos de venta que surte", "Menos de 100", "100–300", "Más de 300"),
    ("Cada cuánto visita cada tienda", "Cada mes", "Cada 2 semanas", "Cada semana"),
    ("Veladoras que ya vende al mes", "No vende o menos de 1,000", "1,000–10,000", "Más de 10,000"),
    ("Vehículos de reparto", "1–2", "3–10", "Más de 10"),
    ("Tamaño del primer pedido", "Menos de $10 mil", "$10–50 mil", "Más de $50 mil"),
], 19):
    for col, t in zip("ABCD", fila):
        celda(di, f"{col}{f}", t); di[f"{col}{f}"].alignment = Alignment(horizontal="left" if col == "A" else "right", wrap_text=True)
encabezado(di, 25, ["Puntaje", "Tipo", "Qué hacer", "", "", ""])
di["C25"].alignment = Alignment(horizontal="left")
for f, (a, b, c) in enumerate([
    ("12 a 15", "A · Grande", "Llamada en menos de 24 horas y visita a su bodega. Ofrecer exclusividad de zona."),
    ("8 a 11", "B · Mediano", "Videollamada en 72 horas. Piloto con 30 a 50 tiendas."),
    ("5 a 7", "C · Chico", "Mensaje con catálogo y pedido mínimo. Sin llamada."),
], 26):
    celda(di, f"A{f}", a); celda(di, f"B{f}", b, negrita=True); celda(di, f"C{f}", c)
celda(di, "A30", "Más de 10,000 veladoras al mes es A directo. Zona con exclusividad ya asignada: lista de espera.", nota=True)
celda(di, "A31", "Las solicitudes del sitio (vela.capitaltorreon.com/distribuir) se guardan ya calificadas con esta tabla.", nota=True)
di.freeze_panes = "B6"

# ── 5. Gráfica del Resumen, en grises ──
res, resultados = wb["Resumen"], wb["Resultados"]
barras = BarChart(); barras.type = "col"; barras.title = "Ingresos y EBITDA por mes"; barras.style = 1
barras.y_axis.title = "Pesos"; barras.x_axis.title = "Mes"; barras.height = 9; barras.width = 22; barras.gapWidth = 40
meses = Reference(resultados, min_col=3, max_col=38, min_row=3)
barras.add_data(Reference(resultados, min_col=2, max_col=38, min_row=20), from_rows=True, titles_from_data=True)
barras.set_categories(meses)
barras.series[0].graphicalProperties.solidFill = G300; barras.series[0].graphicalProperties.line.solidFill = G300
barras.series[0].tx = openpyxl.chart.series.SeriesLabel(v="Ingresos")
linea = LineChart()
linea.add_data(Reference(resultados, min_col=2, max_col=38, min_row=63), from_rows=True, titles_from_data=True)
linea.series[0].graphicalProperties.line.solidFill = NEGRO; linea.series[0].graphicalProperties.line.width = 28000
linea.series[0].smooth = False; linea.series[0].tx = openpyxl.chart.series.SeriesLabel(v="EBITDA")
barras += linea
barras.legend.position = "b"; barras.y_axis.number_format = '\\$#,##0'; barras.y_axis.delete = False; barras.x_axis.delete = False
res.add_chart(barras, "G3")

# ── 6. Riesgos: el del cartucho apunta a su hoja ──
rie = wb["Riesgos"]
for fila in rie.iter_rows(min_row=4):
    if fila[1].value and "artucho" in str(fila[1].value):
        fila[3].value = "Supuestos · sección 17 y hoja Cartucho"

# Pestañas: negras las que se leen, grises las de trabajo
for ws in wb:
    ws.sheet_properties.tabColor = NEGRO if ws.title in ("Inicio", "Resumen", "Cartucho", "Todos_ganan", "Distribuidor") else G300
wb.active = 0
wb.properties.creator = _RLR; wb.properties.title = "La Vela · Modelo de negocio v2"; wb.properties.keywords = f"{_k} {_rev}"
wb.save(DESTINO)
print("guardado", DESTINO)
