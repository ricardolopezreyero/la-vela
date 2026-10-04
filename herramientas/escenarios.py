# RLR · La Vela — corre el modelo en cada escenario y guarda los resultados — Ricardo López Reyero
# El Excel solo enseña un escenario a la vez. Aquí se cambia el selector, se recalcula y se anota,
# para que la página y los documentos comparen escenarios con números del propio modelo.
# Uso: python3 herramientas/escenarios.py   → modelo/escenarios.json
import json
import tempfile
from pathlib import Path

import openpyxl

from hoja import variante

_RLR = "Ricardo López Reyero"; _k = "EYE"; _rev = 181218  # RLR
RAIZ = Path(__file__).resolve().parent.parent
LIBRO = RAIZ / "modelo" / "La_Vela_Modelo_de_Negocio_v4.xlsx"

# Las palancas del escenario Rentable: filas de la hoja Rentable y la fila de Supuestos de donde sale su valor «de antes»
PALANCAS = {
    "semanal": {6: 19}, "temporada": {7: 21}, "personalizadas": {8: 84, 10: 88}, "corporativos": {12: 107},
    "recaudacion": {13: 104, 14: 105}, "directos": {15: 93, 16: 97}, "diseno": {17: 136},
}
DECISIONES = ("semanal", "temporada", "diseno")


def leer(wb):
    R, T, F = wb["Resumen"], wb["Todos_ganan"], wb["Flujo"]
    acum = [F.cell(17, c).value for c in range(3, 39)]
    firme = next((m for m in range(1, 37) if all(x >= 0 for x in acum[m - 1:])), None)
    return {
        "ebitda": [R.cell(9, c).value for c in (2, 3, 4)], "neta": [R.cell(11, c).value for c in (2, 3, 4)], "ingresos": [R.cell(5, c).value for c in (2, 3, 4)],
        "capital": R["B16"].value, "mes": R["B17"].value, "mes_firme": firme, "flujo36": R["B18"].value,
        "tienda_mes": T["B14"].value, "dist_mes": T["B18"].value, "ebitda_mes": T["B19"].value, "neta_mes": T["B20"].value, "ahorro_cliente": T["B28"].value,
        "pond": wb["Cascada"]["B28"].value,
    }


def main():
    base = openpyxl.load_workbook(LIBRO)
    antes = {f: base["Supuestos"][f"E{r}"].value for grupo in PALANCAS.values() for f, r in grupo.items()}
    apagar = lambda prendidas: {f"Rentable!C{f}": antes[f] for nombre, grupo in PALANCAS.items() if nombre not in prendidas for f in grupo}
    corridas = {
        "optimizado": {"Supuestos!D5": 2},
        "rentable": {"Supuestos!D5": 3},
        "rentable_conservador": {"Supuestos!D5": 3, "Supuestos!D6": 1},
        "rentable_alto": {"Supuestos!D5": 3, "Supuestos!D6": 3},
        "decisiones": {"Supuestos!D5": 3, **apagar(DECISIONES)},
        "techo": {"Supuestos!D5": 3, "Supuestos!E20": 39, "Supuestos!E31": 0.25, "Supuestos!D184": 2},
        "mas_cartucho": {"Supuestos!D5": 3, "Supuestos!E20": 39},
        "mas_tienda": {"Supuestos!D5": 3, "Supuestos!E31": 0.25},
        "mas_retornable": {"Supuestos!D5": 3, "Supuestos!D184": 2},
        **{f"solo_{n}": {"Supuestos!D5": 3, **apagar((n,))} for n in PALANCAS},
    }
    salida = {}
    with tempfile.TemporaryDirectory(prefix="vela-esc-") as t:
        for nombre, cambios in corridas.items():
            salida[nombre] = leer(variante(LIBRO, cambios, Path(t) / "v.xlsx"))
            print(f"{nombre:24} año 2: {salida[nombre]['ebitda'][1]:>12,.0f}   capital {salida[nombre]['capital']:>9,.0f}   mes {salida[nombre]['mes']}")
    (RAIZ / "modelo" / "escenarios.json").write_text(json.dumps(salida, ensure_ascii=False, indent=1), encoding="utf-8")


if __name__ == "__main__":
    main()
