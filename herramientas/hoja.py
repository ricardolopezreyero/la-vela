# RLR · La Vela — recalcular un Excel con LibreOffice, sin ventana — Ricardo López Reyero
# openpyxl guarda las fórmulas sin su resultado; LibreOffice las calcula al abrir y vuelve a guardar.
import shutil
import subprocess
import tempfile
from pathlib import Path

import openpyxl

_RLR = "Ricardo López Reyero"  # RLR
SOFFICE = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
ERRORES = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")


def recalcular(ruta):
    """Recalcula en su lugar y falla si quedó algún error de fórmula."""
    ruta = Path(ruta).resolve()
    with tempfile.TemporaryDirectory(prefix="vela-calc-") as t:
        subprocess.run([SOFFICE, "--headless", "--norestore", f"-env:UserInstallation=file://{t}/perfil", "--convert-to", "xlsx", "--outdir", f"{t}/out", str(ruta)],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=240)
        shutil.copy(f"{t}/out/{ruta.name}", ruta)
    wb = openpyxl.load_workbook(ruta, data_only=True)
    malos = [f"{ws.title}!{c.coordinate}" for ws in wb for fila in ws.iter_rows() for c in fila if isinstance(c.value, str) and c.value.strip() in ERRORES]
    if malos:
        raise SystemExit(f"Errores de fórmula en {ruta.name}: {malos[:10]}")
    return wb


def variante(origen, cambios, destino):
    """Copia el libro, cambia celdas ({'Hoja!A1': valor}) y lo recalcula. Devuelve el libro con valores."""
    wb = openpyxl.load_workbook(origen)
    for ref, v in cambios.items():
        hoja, celda = ref.split("!")
        wb[hoja][celda] = v
    wb.save(destino)
    return recalcular(destino)
